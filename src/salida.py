"""Escribe los entregables: Excel del mes, reporte por categoria y copia renombrada de los comprobantes.

Nunca toca los originales: los comprobantes se COPIAN con el nombre nuevo (reversible) y se deja un manifiesto.
"""
import csv
import re
import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLORES = {"GRIS": "D9D9D9", "AMARILLO": "FFFF00", "NARANJA": "FFA500", "BLANCO": None}
CELESTE = "B0E0FF"
ENCABEZADO = PatternFill("solid", fgColor="1F3864")
COLS = ["N°", "Fecha", "Suc. Origen", "Desc. Sucursal", "Cod. Operativo", "Referencia", "Concepto",
        "Importe Pesos", "Saldo Pesos", "DETALLE", "N° Comprobante", "Confianza", "Motivo", "Documentos"]


def _fill(hexa):
    return PatternFill("solid", fgColor=hexa) if hexa else PatternFill(fill_type=None)


def _cabecera(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), ENCABEZADO


def excel_mes(filas: list[dict], ruta: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = ruta.stem
    _cabecera(ws, COLS)
    for f in filas:
        ws.append([f["n"], date.fromisoformat(f["fecha"]), f["suc_origen"], f["desc_sucursal"], f["cod_operativo"],
                   f["referencia"], f["concepto"], f["importe"], f["saldo"], f["detalle"] or None,
                   f["comprobante"] or None, f["confianza"], f["motivo"], ", ".join(f["docs"])])
        r = ws.max_row
        ws.cell(r, 2).number_format = "DD/MM/YYYY"
        ws.cell(r, 8).number_format = '#,##0.00;[Red](#,##0.00)'
        for c in range(1, len(COLS) + 1):
            ws.cell(r, c).fill = _fill(COLORES[f["estado"]])
        if f["detalle"] and f["detalle_origen"] == "deducido":
            ws.cell(r, 10).fill = _fill(CELESTE)
    for col, ancho in zip("ABCDEFGHIJKLMN", (5, 11, 8, 18, 8, 11, 60, 16, 16, 46, 26, 10, 46, 24)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    wb.save(ruta)


def categoria(f: dict) -> str:
    if f["estado"] == "GRIS":
        return "Gris - Impuestos y comisiones bancarias"
    return f["detalle"].split(" - ")[0].strip() if f["detalle"] else "Sin categorizar / sin documento"


def reporte(filas: list[dict], ruta: Path, titulo: str, resumen_extra: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Total de movimientos", len(filas)])
    for e in ("GRIS", "BLANCO", "AMARILLO", "NARANJA"):
        ws.append([f"Filas {e}", sum(f["estado"] == e for f in filas)])
        ws.cell(ws.max_row, 1).fill = _fill(COLORES[e])
    ws.append(["Detalle deducido (celeste, a validar)", sum(bool(f["detalle"]) and f["detalle_origen"] == "deducido" for f in filas)])
    ws.append(["Sin detalle (no gris)", sum(not f["detalle"] and f["estado"] != "GRIS" for f in filas)])
    ws.append(["Confianza baja o media", sum(f["confianza"] != "alta" for f in filas)])
    ws.append([])
    for k, v in resumen_extra.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 24

    tot: dict[str, list] = {}
    for f in filas:
        t = tot.setdefault(categoria(f), [0.0, 0])
        t[0] += f["importe"]
        t[1] += 1
    w2 = wb.create_sheet("Totales por categoría")
    _cabecera(w2, ["Categoría", "Total ($)", "Movimientos"])
    for cat, (imp, n) in sorted(tot.items(), key=lambda kv: -abs(kv[1][0])):
        w2.append([cat, round(imp, 2), n])
        w2.cell(w2.max_row, 2).number_format = '#,##0.00;[Red]-#,##0.00'
    ultima = w2.max_row
    w2.append(["TOTAL", f"=SUM(B2:B{ultima})", f"=SUM(C2:C{ultima})"])
    w2.cell(w2.max_row, 2).number_format = '#,##0.00;[Red]-#,##0.00'
    for c in w2[w2.max_row]:
        c.font = Font(bold=True)
    w2.column_dimensions["A"].width = 44
    w2.column_dimensions["B"].width = 20

    w3 = wb.create_sheet("Para revisar")
    _cabecera(w3, ["N°", "Fecha", "Concepto", "Importe", "Estado", "Detalle", "N° Comprobante", "Confianza", "Motivo"])
    for f in filas:
        if f["estado"] in ("AMARILLO", "NARANJA") or f["confianza"] != "alta" or (f["detalle"] and f["detalle_origen"] == "deducido"):
            w3.append([f["n"], f["fecha"], f["concepto"], f["importe"], f["estado"], f["detalle"],
                       f["comprobante"], f["confianza"], f["motivo"]])
            for c in range(1, 10):
                w3.cell(w3.max_row, c).fill = _fill(COLORES[f["estado"]])
    for col, ancho in zip("ABCDEFGHI", (5, 11, 60, 16, 10, 44, 24, 10, 50)):
        w3.column_dimensions[col].width = ancho
    wb.save(ruta)


def _seguro(t: str) -> str:
    return re.sub(r'[<>:"/\\|?*\n\r]+', " ", t).strip()[:120]


def copiar_renombrados(filas, lecturas, origen: Path, destino: Path) -> list[dict]:
    """Copia cada comprobante con nombre '[movs] - [detalle] - [comprobante].ext'. Los no vinculados van a 'Sin movimientos'."""
    destino.mkdir(parents=True, exist_ok=True)
    sin = destino / "Sin movimientos"
    sin.mkdir(exist_ok=True)
    movs_de: dict[str, list] = {}
    for f in filas:
        for d in f["docs"]:
            movs_de.setdefault(d, []).append(f)
    manifiesto = []
    for doc in sorted(lecturas):
        e = lecturas[doc]["extraccion"]
        ext = Path(doc).suffix
        if doc in movs_de:
            fs = movs_de[doc]
            ns = " y ".join(str(f["n"]) for f in fs) if len(fs) <= 2 else ", ".join(str(f["n"]) for f in fs[:-1]) + " y " + str(fs[-1]["n"])
            nuevo = _seguro(f"{ns} - {fs[0]['detalle'] or e['emisor'] or 'sin detalle'} - {e['numero_comprobante'] or 'sin numero'}") + ext
            carpeta = destino
        else:
            nuevo = _seguro(f"x - {e['emisor'] or e['tipo']} - {e['numero_comprobante'] or 'sin numero'} ({doc})") + ext
            carpeta = sin
        shutil.copy2(origen / doc, carpeta / nuevo)
        manifiesto.append({"original": doc, "nuevo": (carpeta.name + "/" if carpeta is sin else "") + nuevo,
                           "movimientos": ",".join(str(f["n"]) for f in movs_de.get(doc, []))})
    with open(destino / "manifiesto_renombrado.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["original", "nuevo", "movimientos"])
        w.writeheader()
        w.writerows(manifiesto)
    return manifiesto
