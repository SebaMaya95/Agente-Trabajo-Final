"""Arma una planilla para que la persona adjudique SOLO los casos donde el agente y el cierre manual difieren.

Por cada diferencia muestra lo que decidio el agente, lo que dice el cierre manual y los documentos involucrados
(con su nombre original, que es privado). La persona completa "Quien tiene razon" y un comentario.
Salida privada: datos_privados/<mes>/revision_desacuerdos.xlsx  (contiene datos reales: no va al repositorio).
Uso: python planilla_desacuerdos.py <carpeta_corrida_privada> <mes>
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from evaluar import cargar_esperado, cargar_mapa  # noqa: E402


def main(run: Path, mes: str):
    S = json.loads((run / "salida.json").read_text(encoding="utf-8"))["filas"]
    L = json.loads((run / "lecturas.json").read_text(encoding="utf-8"))
    R = json.loads((run / "evaluacion.json").read_text(encoding="utf-8"))
    esp = cargar_esperado(mes)
    mapa = cargar_mapa(mes)
    nombres = {r["doc"]: r["nombres_originales"] for r in csv.DictReader(open(RAIZ / f"datos_privados/{mes}/esperado/mapa_documentos.csv", encoding="utf-8-sig"))}
    fila = {f["n"]: f for f in S}

    def doc_txt(d):
        e = L[d]["extraccion"]
        imp = ", ".join(f"{i['etiqueta']} {i['monto']:,.2f}" for i in e["importes"][:3])
        return f"{d} | archivo original: {nombres.get(d, '?')[:120]} | {e['tipo']} · {e['emisor'] or 's/emisor'} · N° {e['numero_comprobante'] or 's/n'} · {imp}"

    casos: dict = {}
    for er in R["errores"]:
        c = casos.setdefault(("mov", er["n"]), {"tipos": set()})
        c["tipos"].add("Comprobante distinto" if er["campo"] == "comprobante" else "Detalle distinto")
    for d in R["documentos"]["sin_movimiento_vinculados_de_mas"]:
        for n in [f["n"] for f in S if d in f["docs"]]:
            casos.setdefault(("mov", n), {"tipos": set()})["tipos"].add("Documento vinculado que el cierre manual no vincula")
    for d, movs in mapa.items():
        if not movs:
            continue
        vinc = {f["n"] for f in S if d in f["docs"]}
        if vinc != movs:
            casos[("doc", d)] = {"tipos": {"Documento con movimiento esperado que el agente no vinculó (o vinculó a medias)"}, "movs": movs, "vinc": vinc}

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Cómo completar"
    for l in ["Revisión de desacuerdos entre el agente (v5) y tu cierre manual — mes " + mes, "",
              "1. Solo están los casos donde hay diferencia. Los demás coinciden.",
              "2. Para cada fila, mirá lo que decidió el agente, lo que dice tu cierre manual y los documentos (con su archivo original).",
              "3. En la columna 'Quién tiene razón' elegí: Agente / Cierre manual / Ambos válidos / No sé.",
              "   - 'Agente': tu cierre manual estaba equivocado o incompleto en ese punto.",
              "   - 'Cierre manual': el agente se equivocó.",
              "   - 'Ambos válidos': es una diferencia de criterio o de forma (por ejemplo, otra redacción del Detalle).",
              "4. Si querés, dejá un comentario breve (por qué, o qué regla faltó).",
              "", "Importante: este archivo tiene datos reales del cliente. No lo subas a GitHub."]:
        ws0.append([l])
    ws0.column_dimensions["A"].width = 120
    ws0["A1"].font = Font(bold=True, size=13)

    ws = wb.create_sheet("Diferencias")
    cab = ["Caso", "Tipo de diferencia", "Fecha", "Concepto del banco", "Importe", "AGENTE: comprobante", "AGENTE: detalle",
           "AGENTE: confianza y motivo", "CIERRE MANUAL: comprobante", "CIERRE MANUAL: detalle", "Documentos involucrados",
           "Quién tiene razón", "Comentario"]
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"Agente,Cierre manual,Ambos válidos,No sé"', allow_blank=True)
    ws.add_data_validation(dv)

    def clave(k):
        return (k[1] if k[0] == "mov" else min(casos[k]["movs"]), k[0])
    for k in sorted(casos, key=clave):
        c = casos[k]
        if k[0] == "mov":
            n = k[1]
            f, e = fila[n], esp[n]
            docs = "\n".join(doc_txt(d) for d in f["docs"]) or "(el agente no vinculó ningún documento)"
            docs_esp = [d for d, m in mapa.items() if n in m]
            if docs_esp:
                docs += "\n— Según el cierre manual el documento sería: " + "; ".join(nombres.get(d, "?")[:90] for d in docs_esp)
            fila_xl = [f"Mov. {n}", " + ".join(sorted(c["tipos"])), f["fecha"], f["concepto"], f["importe"], f["comprobante"], f["detalle"],
                       f"{f['confianza']} — {f['motivo']}", e["comprobante"], e["detalle"] + ("  (en celeste: sin validar)" if e["celeste"] else ""), docs]
        else:
            d = k[1]
            ns = sorted(c["movs"])
            fila_xl = [f"Doc. {d}", " + ".join(sorted(c["tipos"])), fila[ns[0]]["fecha"], " / ".join(f"Mov. {n}: {fila[n]['concepto'][:60]}" for n in ns),
                       sum(fila[n]["importe"] for n in ns), " ; ".join(fila[n]["comprobante"] or "(vacío)" for n in ns), "",
                       f"El agente lo vinculó a: {sorted(c['vinc']) or 'ningún movimiento'}", " ; ".join(esp[n]["comprobante"] or "(vacío)" for n in ns),
                       "", doc_txt(d)]
        ws.append(fila_xl + ["", ""])
        r = ws.max_row
        ws.cell(r, 5).number_format = '#,##0.00;[Red]-#,##0.00'
        dv.add(ws.cell(r, 12))
        for col in range(1, 14):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLM", (10, 26, 11, 44, 15, 20, 30, 40, 20, 30, 70, 16, 34)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    salida = RAIZ / f"datos_privados/{mes}/revision_desacuerdos.xlsx"
    wb.save(salida)
    print(f"{len(casos)} casos con diferencia -> {salida}")
    from collections import Counter
    print(Counter(t for c in casos.values() for t in c["tipos"]))


if __name__ == "__main__":
    main(Path(sys.argv[1]), sys.argv[2])
